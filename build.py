#!/usr/bin/env python3
"""Build the VADO technical and financial proposal.

    python3 build.py                          # blank template, tags intact
    python3 build.py --data jobs/mkm.json     # a filled proposal
    python3 build.py --out out/TFP.pptx

The output is a plain .pptx. No macros, no Excel, no Windows.
"""
from __future__ import annotations

import argparse
import importlib
import json
import os
import sys

from pptx import Presentation
from pptx.oxml.ns import qn

from vado.context import ROOT, Ctx
from vado.deck import pt

SLIDE_MODULES = [
    'vado.slides.s01_cover',
    'vado.slides.s02_contents',
    'vado.slides.s03_intro',
    'vado.slides.s04_project',
    'vado.slides.s05_spatial',
    'vado.slides.s06_scope',
    'vado.slides.s08_services',
    'vado.slides.s09_phases',
    'vado.slides.s10_method',
    'vado.slides.s11_timeline',
    'vado.slides.s12_financial',
    'vado.slides.s13_agreement',
    'vado.slides.s14_contact',
    'vado.slides.s15_closing',
]

BLANK_LAYOUT = 6


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--theme', default=os.path.join(ROOT, 'theme.json'))
    ap.add_argument('--content', default=os.path.join(ROOT, 'content.json'))
    ap.add_argument('--data', help='JSON of merge-field values to overlay')
    ap.add_argument('--out', default=os.path.join(ROOT, 'out', 'VADO-TFP.pptx'))
    ap.add_argument('--only', type=int, nargs='*',
                    help='build only these slide numbers, for quick checks')
    args = ap.parse_args()

    ctx = Ctx(args.theme, args.content, args.data)
    problems = validate(ctx)
    if problems:
        for p in problems:
            print('ERROR: ' + p, file=sys.stderr)
        return 1

    prs = Presentation()
    prs.slide_width = pt(ctx.geo('canvas', 'width_pt'))
    prs.slide_height = pt(ctx.geo('canvas', 'height_pt'))
    set_master_background(prs, ctx.asset('master-bg.png'))

    layout = prs.slide_layouts[BLANK_LAYOUT]
    wanted = set(args.only) if args.only else None

    # A module may declare pages(ctx) and take a page index. The scope does,
    # because a project with more disciplines needs more room and the deck
    # should grow rather than the text running into the footer.
    manifest, number = [], 0
    for name in SLIDE_MODULES:
        mod = importlib.import_module(name)
        count = mod.pages(ctx) if hasattr(mod, 'pages') else 1
        for page in range(count):
            number += 1
            manifest.append({'slide': number, 'module': name.rsplit('.', 1)[-1],
                             'title': getattr(mod, 'TITLE', ''), 'page': page})
            if wanted and number not in wanted:
                continue
            slide = prs.slides.add_slide(layout)
            if hasattr(mod, 'pages'):
                mod.build(ctx, slide, page)
            else:
                mod.build(ctx, slide)

    if ctx.unknown_tags:
        print('WARNING: no value for ' + ', '.join(sorted(ctx.unknown_tags)),
              file=sys.stderr)

    os.makedirs(os.path.dirname(args.out), exist_ok=True)
    prs.save(args.out)
    # The deck no longer has a fixed slide count, so anything that wants to
    # talk about "slide 7" needs to be told what is on it.
    with open(args.out + '.manifest.json', 'w', encoding='utf-8') as fh:
        json.dump(manifest, fh, indent=2)
    print(f'wrote {args.out}  ({len(prs.slides.__iter__.__self__._sldIdLst)} slides)')
    return 0


def validate(ctx):
    """Catch the mistakes that used to reach the client."""
    out = []

    pcts = []
    for key in ('p1', 'p2', 'p3', 'p4'):
        raw = str(ctx.field(key)).strip().rstrip('%')
        if raw.startswith('{{'):
            pcts = None
            break
        try:
            v = float(raw)
        except ValueError:
            out.append(f'payment {key} is not a number: {ctx.field(key)!r}')
            pcts = None
            break
        pcts.append(v if v > 1 else v * 100)
    if pcts and abs(sum(pcts) - 100.0) > 0.01:
        out.append(f'payments must total 100%, they total {sum(pcts):.1f}%')

    keys = [s['key'] for s in ctx.sections]
    if len(keys) != len(set(keys)):
        out.append('content.json sections has a duplicate key')
    if len(keys) != ctx.geo('chrome', 'rail')['count']:
        out.append(f'the side rail expects {ctx.geo("chrome", "rail")["count"]} '
                   f'sections, content.json has {len(keys)}')
    return out


def set_master_background(prs, image_path):
    """Stretch the paper texture across the slide master, once."""
    master = prs.slide_masters[0]
    image_part, rId = master.part.get_or_add_image_part(image_path)
    cSld = master._element.find(qn('p:cSld'))

    old = cSld.find(qn('p:bg'))
    if old is not None:
        cSld.remove(old)

    bg = cSld.makeelement(qn('p:bg'), {})
    bgPr = bg.makeelement(qn('p:bgPr'), {})
    blipFill = bgPr.makeelement(qn('a:blipFill'), {'dpi': '0', 'rotWithShape': '1'})
    blip = blipFill.makeelement(qn('a:blip'), {qn('r:embed'): rId})
    blipFill.append(blip)
    blipFill.append(blipFill.makeelement(qn('a:srcRect'), {}))
    stretch = blipFill.makeelement(qn('a:stretch'), {})
    stretch.append(stretch.makeelement(qn('a:fillRect'), {}))
    blipFill.append(stretch)
    bgPr.append(blipFill)
    bgPr.append(bgPr.makeelement(qn('a:effectLst'), {}))
    bg.append(bgPr)
    cSld.insert(0, bg)


if __name__ == '__main__':
    raise SystemExit(main())
