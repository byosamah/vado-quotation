#!/usr/bin/env python3
"""Build everything the browser self-check needs, then say what to paste.

    python3 html/tests/prepare.py [jobs/<id>.json]

The check itself runs in the browser, because the questions it answers are
browser questions: what the layout engine did with a paragraph, where a
table actually ended, which way round the bidi algorithm drew a number.
This script prepares the fixtures it compares against: the text of the
PowerPoint built from the same data, which is the thing the HTML has to
agree with.
"""
from __future__ import annotations

import json
import os
import shutil
import subprocess
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, '..', '..'))
OUT = os.path.join(ROOT, 'out')
PY = sys.executable


def pptx_text(path):
    from pptx import Presentation

    def walk(shape, acc):
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                s = ''.join(r.text for r in p.runs).strip()
                if s:
                    acc.append(s)
        if getattr(shape, 'has_table', False) and shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    for p in cell.text_frame.paragraphs:
                        s = ''.join(r.text for r in p.runs).strip()
                        if s:
                            acc.append(s)
        if shape.shape_type == 6:          # a group
            for sub in shape.shapes:
                walk(sub, acc)

    pages = []
    for slide in Presentation(path).slides:
        acc = []
        for sh in slide.shapes:
            walk(sh, acc)
        pages.append(acc)
    return pages


def main():
    job = sys.argv[1] if len(sys.argv) > 1 else None
    os.makedirs(OUT, exist_ok=True)

    subprocess.check_call([PY, os.path.join(ROOT, 'html', 'build_html.py'),
                           '--out', os.path.join(OUT, 'VADO-TFP.html')])

    targets = [(None, '_pptx_text.json', '_cmp.pptx')]
    if job:
        targets.append((job, '_pptx_text_job.json', '_cmpjob.pptx'))

    for data, textfile, pptx in targets:
        cmd = [PY, os.path.join(ROOT, 'build.py'), '--out', os.path.join(OUT, pptx)]
        if data:
            cmd += ['--data', data]
        subprocess.check_call(cmd, stdout=subprocess.DEVNULL)
        pages = pptx_text(os.path.join(OUT, pptx))
        with open(os.path.join(OUT, textfile), 'w', encoding='utf-8') as fh:
            json.dump(pages, fh, ensure_ascii=False)
        print(f'  {textfile}: {len(pages)} slides, '
              f'{sum(len(p) for p in pages)} paragraphs')

    shutil.copy(os.path.join(HERE, 'selfcheck.js'), os.path.join(OUT, '_check.js'))
    if job:
        shutil.copy(job, os.path.join(OUT, '_job.json'))

    print(f"""
Now serve out/ and open the file:

    cd {OUT} && python3 -m http.server 8756 --bind 127.0.0.1
    open http://127.0.0.1:8756/VADO-TFP.html

Then paste this into the browser console:

    eval(await (await fetch('/_check.js')).text())
    await vadoSelfCheck('/_pptx_text.json')          // the blank template
""")
    if job:
        print("""    VADO_APP.ctx.job = await fetch('/_job.json').then(r=>r.json())
    VADO_APP.changed()
    await vadoSelfCheck('/_pptx_text_job.json')      // the filled proposal
""")
    print('`passed` must be true and `fail` must be empty for both.\n'
          'A local web server is needed because the extension and fetch() '
          'will not read file:// URLs. Opening the .html by hand does not '
          'need one.\n')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
